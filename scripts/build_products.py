#!/usr/bin/env python3
"""Build js/products.js from 2026-09-10.xlsx + 2026-09-11.xlsx."""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX_MAIN = ROOT / "2026-09-10.xlsx"
XLSX_EXTRA = ROOT / "2026-09-11.xlsx"
OUT = ROOT / "js" / "products.js"
IMG_DIR = ROOT / "img" / "products"
CNY_USD = 7.2

CAT_MAP = {
    "Shoes": "shoes",
    "Jordan 4": "shoes",
    "Air force": "shoes",
    "Nike Shox": "shoes",
    "Dior b30": "shoes",
    "Dior B30": "shoes",
    "Tops": "t-shirts",
    "Polo": "t-shirts",
    "Shirts": "t-shirts",
    "Hoodies": "hoodies",
    "Nike Tech": "hoodies",
    "Jacket": "jackets",
    "Bottoms": "pants",
    "Pants": "pants",
    "Shorts": "pants",
    "Underwear": "underwear",
    "Bag": "bags",
    "Accessories": "accessories",
    "Watch": "watches",
    "Electronics": "other",
    "Airpods": "other",
    "Cap": "headwear",
    "Hat": "headwear",
    "Jersey": "jersey",
    "Set": "sets",
    "Tracksuit": "sets",
    "Suit": "sets",
    "Sunglasses": "glasses",
    "Dress": "other",
    "Backpack": "bags",
    "Belt": "accessories",
    "Socks": "accessories",
    "Jeans": "pants",
    "Sweater": "hoodies",
    "Down Jacket": "jackets",
    "Vest": "jackets",
    "Perfume": "perfume",
}

CAT_LABEL = {
    "shoes": "Shoes",
    "t-shirts": "T-shirts",
    "hoodies": "Hoodies",
    "jackets": "Jackets",
    "pants": "Pants",
    "underwear": "Underwear",
    "bags": "Bags",
    "accessories": "Accessories",
    "watches": "Watches",
    "headwear": "Headwear",
    "jersey": "Jersey",
    "sets": "Sets",
    "glasses": "Glasses",
    "perfume": "Perfume",
    "bricks": "Bricks",
    "other": "Other",
}

META_TAGS = {
    "hot nike",
    "top trending",
    "new arrivals",
    "hot picks",
    "scarce lv",
    "streetwear",
    "football",
}

TAG_CAT = {
    "shoes": "shoes",
    "slides": "shoes",
    "sneakers": "shoes",
    "jordan 4": "shoes",
    "t-shirt": "t-shirts",
    "shirts": "t-shirts",
    "shirt": "t-shirts",
    "polo": "t-shirts",
    "hoodies": "hoodies",
    "hoodie": "hoodies",
    "sweater": "hoodies",
    "jacket": "jackets",
    "down jacket": "jackets",
    "vest": "jackets",
    "shorts": "pants",
    "pants": "pants",
    "jeans": "pants",
    "bottoms": "pants",
    "bag": "bags",
    "backpack": "bags",
    "accessories": "accessories",
    "cap": "headwear",
    "hat": "headwear",
    "belt": "accessories",
    "socks": "accessories",
    "sunglasses": "glasses",
    "bracelet": "accessories",
    "necklace": "accessories",
    "watch": "watches",
    "electronics": "other",
    "underwear": "underwear",
    "nike tech": "hoodies",
    "set": "sets",
    "tracksuit": "sets",
    "suit": "sets",
    "jersey": "jersey",
    "dress": "other",
    "perfume": "perfume",
}


def item_id_from_url(url: str) -> str:
    if not url:
        return ""
    m = re.search(r"itemID=(\d+)", str(url), re.I)
    if m:
        return m.group(1)
    m = re.search(r"(\d{8,})", str(url))
    return m.group(1) if m else ""


def guess_cat(raw, title: str) -> str:
    t = (title or "").lower()
    bag = any(k in t for k in ("bag", "wallet", "shoulder", "backpack", "包"))
    if re.search(r"\b(lego|bricks?)\b", t) or "积木" in t or "building block" in t:
        return "bricks"
    if (re.search(r"\b(sunglasses?|eyewear|ray-?ban|oakley)\b", t) or "眼镜" in t) and "jacket" not in t:
        return "glasses"
    if (re.search(r"\b(watch(?:es)?|rolex|patek|omega|casio|swatch)\b", t) or "手表" in t) and not bag:
        return "watches"
    if (re.search(r"\b(perfume|cologne|parfum)\b", t) or "香水" in t) and not bag:
        return "perfume"
    if any(k in t for k in ("shoe", "slipper", "slide", "sneaker", "jordan", "拖鞋", "鞋")):
        return "shoes"
    if any(k in t for k in ("jersey", "football kit", "soccer", "球衣")):
        return "jersey"
    if any(k in t for k in ("hoodie", "sweater", "fleece", "卫衣", "帽衫")):
        return "hoodies"
    if "帽衫" not in t and any(k in t for k in ("cap", "hat", "beanie", "trucker", "帽")):
        return "headwear"
    if any(k in t for k in ("tracksuit", "套装")) or re.search(r"\bset\b", t):
        return "sets"
    if any(k in t for k in ("jacket", "coat", "windbreaker", "bomber", "夹克", "外套")):
        return "jackets"
    if raw in CAT_MAP:
        return CAT_MAP[raw]
    if any(k in t for k in ("short", "pant", "jean", "裤")):
        return "pants"
    if any(k in t for k in ("tee", "t-shirt", "t恤", "shirt", "polo")):
        return "t-shirts"
    if any(k in t for k in ("bag", "wallet", "包", "belt", "sock", "bracelet")):
        return "bags" if bag else "accessories"
    return "other"


def parse_tags(raw: str):
    parts = [p.strip() for p in str(raw or "").replace("\n", ",").split(",") if p.strip()]
    brand = ""
    cat = ""
    tags = []
    for part in parts:
        key = part.lower()
        tags.append(part)
        if key in META_TAGS:
            continue
        if key in TAG_CAT:
            if not cat:
                cat = TAG_CAT[key]
            continue
        if not brand:
            brand = part
    return brand, cat, tags


def usd(cny) -> float:
    try:
        v = float(cny or 0)
    except (TypeError, ValueError):
        return 0.0
    return round(v / CNY_USD, 2) if v else 0.0


def image_for(item_id: str) -> str:
    path = IMG_DIR / f"{item_id}.webp"
    if path.exists() and path.stat().st_size > 800:
        return f"img/products/{item_id}.webp"
    return ""


def load_main() -> dict[str, dict]:
    wb = openpyxl.load_workbook(XLSX_MAIN, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = {name: i for i, name in enumerate(rows[0])}
    out: dict[str, dict] = {}
    for r in rows[1:]:
        if not r or not r[h["商品id"]]:
            continue
        iid = str(r[h["商品id"]]).strip()
        if not iid or not iid.isdigit():
            continue
        title = str(r[h["展示商品标题"]] or r[h["商品标题"]] or "Find").strip()
        brand = str(r[h["品牌"]] or "").strip()
        raw_cat = r[h["品牌分类"]] or ""
        cat = guess_cat(raw_cat, title)
        opens = float(r[h["打开商品详情页数"]] or 0)
        url = str(r[h["商品url"]] or f"https://weidian.com/item.html?itemID={iid}").strip()
        out[iid] = {
            "id": iid,
            "title": title[:100],
            "brand": brand,
            "category": cat,
            "categoryLabel": CAT_LABEL.get(cat, "Other"),
            "price": usd(r[h["商品优惠价"]]),
            "url": url,
            "image": image_for(iid),
            "shop": str(r[h["卖家店铺名称"]] or "").strip(),
            "tags": [t for t in [brand, CAT_LABEL.get(cat, "")] if t],
            "opens": opens,
            "featured": opens > 200,
            "hot": opens > 800,
            "qc": True,
            "source": "2026-09-10",
        }
    return out


def load_extra(existing: dict[str, dict]) -> int:
    wb = openpyxl.load_workbook(XLSX_EXTRA, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    added = 0
    for r in rows[1:]:
        if not r:
            continue
        title = str(r[0] or "").strip()
        url = str(r[1] or "").strip()
        tags_raw = str(r[2] or "").strip() if len(r) > 2 else ""
        iid = item_id_from_url(url)
        if not iid or not title:
            continue
        brand, cat, tags = parse_tags(tags_raw)
        if iid in existing:
            # Enrich tags / hot flags from 09-11 lists
            p = existing[iid]
            for t in tags:
                if t not in p["tags"]:
                    p["tags"].append(t)
            if any(t.lower() in ("top trending", "hot picks", "hot nike") for t in tags):
                p["hot"] = True
                p["featured"] = True
            continue
        if not cat:
            cat = guess_cat("", title)
        existing[iid] = {
            "id": iid,
            "title": title[:100],
            "brand": brand,
            "category": cat,
            "categoryLabel": CAT_LABEL.get(cat, "Other"),
            "price": 0,
            "url": url or f"https://weidian.com/item.html?itemID={iid}",
            "image": image_for(iid),
            "shop": "",
            "tags": tags or ([brand, CAT_LABEL.get(cat, "")] if brand else [CAT_LABEL.get(cat, "Other")]),
            "opens": 0,
            "featured": any(t.lower() in ("top trending", "hot picks") for t in tags),
            "hot": any(t.lower() in ("top trending", "hot nike", "hot picks") for t in tags),
            "qc": True,
            "source": "2026-09-11",
        }
        added += 1
    return added


def main() -> None:
    products_map = load_main()
    added = load_extra(products_map)
    products = sorted(products_map.values(), key=lambda p: (-p["opens"], p["title"].lower()))
    cats = sorted({p["category"] for p in products})
    with_img = sum(1 for p in products if p["image"])
    payload = {
        "updated": "23 Sep 2026",
        "count": len(products),
        "categories": [{"slug": c, "label": CAT_LABEL.get(c, c.title())} for c in cats],
        "products": products,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        "window.KakoHub = window.KakoHub || {};\n"
        "KakoHub.catalog = "
        + json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
        + ";\n",
        encoding="utf-8",
    )
    print(
        f"wrote {OUT.name}: {len(products)} products "
        f"(+{added} from 09-11), {with_img} with images, {OUT.stat().st_size // 1024} KB"
    )


if __name__ == "__main__":
    main()
